import os
import sys
import random
import openpyxl
import django

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
django.setup()

from core.models import Estudiante, InformeValorativo

DIRECTORES = {
    "601": "CARDENAS FORERO MARTHA",
    "602": "GARCIA HERNANDEZ PEDRO",
    "603": "LOPEZ DIAZ GLORIA",
    "701": "SUAREZ BLANCO CARLOS",
    "702": "HERRERA MORENO PATRICIA",
    "703": "RODRIGUEZ BERNAL JAVIER",
    "801": "GUTIERREZ SANCHEZ DIANA",
    "802": "RAMIREZ ORTIZ FERNANDO",
    "803": "CASTRO CELIS MONICA",
    "901": "MARTINEZ DUARTE ALVARO",
    "902": "VARGAS PINZON LILIANA",
    "903": "CALDERON CAMACHO RAUL",
    "1001": "VELASQUEZ GAITAN SANDRA",
    "1002": "MORALES CASTILLO HECTOR",
    "1101": "ROJAS GUTIERREZ LILIANA ASTRID",
    "1102": "BUITRAGO GUERRERO ANDRES",
}

GRADOS = {
    "601": "SEXTO", "602": "SEXTO", "603": "SEXTO",
    "701": "SEPTIMO", "702": "SEPTIMO", "703": "SEPTIMO",
    "801": "OCTAVO", "802": "OCTAVO", "803": "OCTAVO",
    "901": "NOVENO", "902": "NOVENO", "903": "NOVENO",
    "1001": "DECIMO", "1002": "DECIMO",
    "1101": "UNDECIMO", "1102": "UNDECIMO",
}


def generar_detalles_etapa(desempeno, seed_val):
    rnd = random.Random(seed_val)
    if desempeno in ("SUPERIOR", "S"):
        des = "S"
        p = rnd.randint(9, 10)
        rd = rnd.randint(18, 20)
        r = rnd.randint(28, 30)
        a = rnd.randint(27, 30)
        pt = rnd.randint(9, 10)
    elif desempeno in ("ALTO", "A"):
        des = "A"
        p = rnd.randint(8, 9)
        rd = rnd.randint(16, 18)
        r = rnd.randint(24, 27)
        a = rnd.randint(24, 26)
        pt = rnd.randint(8, 9)
    elif desempeno in ("BÁSICO", "BASICO", "BS"):
        des = "BS"
        p = rnd.randint(7, 8)
        rd = rnd.randint(14, 15)
        r = rnd.randint(21, 23)
        a = rnd.randint(21, 23)
        pt = rnd.randint(7, 8)
    else:  # BAJO
        des = "BJ"
        p = rnd.randint(5, 6)
        rd = rnd.randint(10, 13)
        r = rnd.randint(15, 19)
        a = rnd.randint(15, 19)
        pt = rnd.randint(5, 6)
    total = p + rd + r + a + pt
    return {"p": p, "rd": rd, "r": r, "a": a, "pt": pt, "total": total, "des": des}


def variante_desempeno(des_actual, rnd):
    if des_actual == "SUPERIOR":
        return rnd.choice(["SUPERIOR", "SUPERIOR", "ALTO"])
    elif des_actual == "ALTO":
        return rnd.choice(["SUPERIOR", "ALTO", "ALTO", "BÁSICO"])
    elif des_actual == "BÁSICO":
        return rnd.choice(["ALTO", "BÁSICO", "BÁSICO", "BAJO"])
    else:
        return rnd.choice(["BÁSICO", "BAJO", "BAJO"])


def importar():
    ruta_excel = os.path.join(BASE_DIR, "sabana_notas_ceis_2026.xlsx")
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)

    estudiantes_objs = []
    datos_por_doc = {}

    print(f"Leyendo {ruta_excel}...")

    for sheetname in wb.sheetnames:
        if sheetname not in GRADOS:
            continue
        ws = wb[sheetname]
        headers = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

        col_doc = 1
        col_nom = 2

        materias_cols = {}
        for idx in range(3, len(headers)):
            h = headers[idx]
            if h in ("Desempeño General", "Estado Académico"):
                break
            materias_cols[h] = idx

        es_media = sheetname in ["1001", "1002", "1101", "1102"]

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[col_doc] is None or row[col_nom] is None:
                continue

            doc_raw = str(row[col_doc]).strip()
            if not doc_raw:
                continue
            try:
                documento = str(int(float(doc_raw)))
            except Exception:
                documento = doc_raw

            if documento in datos_por_doc:
                continue

            nombre_completo = str(row[col_nom]).strip().upper()
            partes = nombre_completo.split()
            if len(partes) >= 3:
                apellidos = f"{partes[0]} {partes[1]}"
                nombres = " ".join(partes[2:])
            elif len(partes) == 2:
                apellidos = partes[0]
                nombres = partes[1]
            else:
                apellidos = nombre_completo
                nombres = ""

            rnd = random.Random(int(documento) if documento.isdigit() else 12345)

            estudiantes_objs.append(
                Estudiante(
                    documento=documento,
                    tipo_documento="Tarjeta de Identidad",
                    nombres=nombres,
                    apellidos=apellidos,
                    grado=GRADOS[sheetname],
                    curso=sheetname,
                    director_grupo=DIRECTORES.get(sheetname, "ROJAS GUTIERREZ LILIANA ASTRID"),
                    acudiente=f"Familia {apellidos}",
                    telefono_acudiente="310" + "".join([str(rnd.randint(0, 9)) for _ in range(7)]),
                    activo=True,
                )
            )

            areas = []
            if es_media:
                asigs_cientifico = []
                for mat_nombre in ["Química", "Física"]:
                    if mat_nombre in materias_cols:
                        nota_val = row[materias_cols[mat_nombre]] or "BÁSICO"
                        asigs_cientifico.append({
                            "nombre": mat_nombre.upper(),
                            "p1": generar_detalles_etapa(variante_desempeno(nota_val, rnd), rnd.randint(1, 10000)),
                            "p2": generar_detalles_etapa(variante_desempeno(nota_val, rnd), rnd.randint(1, 10000)),
                            "p3": generar_detalles_etapa(nota_val, rnd.randint(1, 10000)),
                            "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                        })
                asigs_cientifico.append({
                    "nombre": "BIOLOGÍA",
                    "p1": generar_detalles_etapa("ALTO", rnd.randint(1, 10000)),
                    "p2": generar_detalles_etapa("BÁSICO", rnd.randint(1, 10000)),
                    "p3": generar_detalles_etapa("ALTO", rnd.randint(1, 10000)),
                    "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                })
                areas.append({"nombre": "PENSAMIENTO CIENTÍFICO", "asignaturas": asigs_cientifico})

                nota_mat = row[materias_cols.get("Matemáticas", 3)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO MATEMÁTICO",
                    "asignaturas": [{
                        "nombre": "MATEMÁTICAS",
                        "p1": generar_detalles_etapa(variante_desempeno(nota_mat, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(nota_mat, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(nota_mat, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                asigs_com = []
                for mat_nombre in ["Lengua Castellana", "Inglés"]:
                    if mat_nombre in materias_cols:
                        n = row[materias_cols[mat_nombre]] or "ALTO"
                        asigs_com.append({
                            "nombre": mat_nombre.upper(),
                            "p1": generar_detalles_etapa(variante_desempeno(n, rnd), rnd.randint(1, 10000)),
                            "p2": generar_detalles_etapa(variante_desempeno(n, rnd), rnd.randint(1, 10000)),
                            "p3": generar_detalles_etapa(n, rnd.randint(1, 10000)),
                            "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                        })
                areas.append({"nombre": "COMUNICACIONES (HUMANIDADES)", "asignaturas": asigs_com})

                n_pol = row[materias_cols.get("Política y Economía", 7)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO SOCIAL",
                    "asignaturas": [
                        {
                            "nombre": "P. POLITICO Y ECONOMICO",
                            "p1": generar_detalles_etapa(variante_desempeno(n_pol, rnd), rnd.randint(1, 10000)),
                            "p2": generar_detalles_etapa(variante_desempeno(n_pol, rnd), rnd.randint(1, 10000)),
                            "p3": generar_detalles_etapa(n_pol, rnd.randint(1, 10000)),
                            "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                        },
                        {
                            "nombre": "PENSAMIENTO SOCIAL",
                            "p1": generar_detalles_etapa("BÁSICO", rnd.randint(1, 10000)),
                            "p2": generar_detalles_etapa("ALTO", rnd.randint(1, 10000)),
                            "p3": generar_detalles_etapa("BÁSICO", rnd.randint(1, 10000)),
                            "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                        }
                    ]
                })

                n_fil = row[materias_cols.get("Filosofía", 8)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO FILOSOFICO",
                    "asignaturas": [{
                        "nombre": "FILOSOFIA",
                        "p1": generar_detalles_etapa(variante_desempeno(n_fil, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_fil, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_fil, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                n_ef = row[materias_cols.get("Educación Física", 10)] or "SUPERIOR"
                areas.append({
                    "nombre": "CULTURA FÍSICA Y DEPORTIVA",
                    "asignaturas": [{
                        "nombre": "EDUCACIÓN FÍSICA",
                        "p1": generar_detalles_etapa(variante_desempeno(n_ef, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_ef, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_ef, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                n_art = row[materias_cols.get("Educación Artística", 11)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO ARTÍSTICO",
                    "asignaturas": [{
                        "nombre": "EDUCACIÓN ARTÍSTICA",
                        "p1": generar_detalles_etapa(variante_desempeno(n_art, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_art, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_art, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                n_et = row[materias_cols.get("Ética y Valores", 12)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO ÉTICO Y CIUDADANO",
                    "asignaturas": [{
                        "nombre": "EDUCACIÓN ÉTICA Y VALORES",
                        "p1": generar_detalles_etapa(variante_desempeno(n_et, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_et, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_et, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                n_tec = row[materias_cols.get("Tecnología e Informática", 9)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO TECNOLÓGICO",
                    "asignaturas": [{
                        "nombre": "TECNOLOGÍA E INFORMÁTICA",
                        "p1": generar_detalles_etapa(variante_desempeno(n_tec, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_tec, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_tec, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })
            else:
                n_nat = row[materias_cols.get("Ciencias Naturales", 5)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO CIENTÍFICO",
                    "asignaturas": [{
                        "nombre": "BIOLOGÍA / CIENCIAS NATURALES",
                        "p1": generar_detalles_etapa(variante_desempeno(n_nat, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_nat, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_nat, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                n_mat = row[materias_cols.get("Matemáticas", 3)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO MATEMÁTICO",
                    "asignaturas": [{
                        "nombre": "MATEMÁTICAS",
                        "p1": generar_detalles_etapa(variante_desempeno(n_mat, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_mat, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_mat, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                asigs_com = []
                for mat_nombre in ["Lengua Castellana", "Inglés"]:
                    if mat_nombre in materias_cols:
                        n = row[materias_cols[mat_nombre]] or "ALTO"
                        asigs_com.append({
                            "nombre": mat_nombre.upper(),
                            "p1": generar_detalles_etapa(variante_desempeno(n, rnd), rnd.randint(1, 10000)),
                            "p2": generar_detalles_etapa(variante_desempeno(n, rnd), rnd.randint(1, 10000)),
                            "p3": generar_detalles_etapa(n, rnd.randint(1, 10000)),
                            "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                        })
                areas.append({"nombre": "COMUNICACIONES (HUMANIDADES)", "asignaturas": asigs_com})

                n_soc = row[materias_cols.get("Ciencias Sociales", 6)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO SOCIAL",
                    "asignaturas": [{
                        "nombre": "CIENCIAS SOCIALES",
                        "p1": generar_detalles_etapa(variante_desempeno(n_soc, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_soc, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_soc, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                n_ef = row[materias_cols.get("Educación Física", 9)] or "SUPERIOR"
                areas.append({
                    "nombre": "CULTURA FÍSICA Y DEPORTIVA",
                    "asignaturas": [{
                        "nombre": "EDUCACIÓN FÍSICA",
                        "p1": generar_detalles_etapa(variante_desempeno(n_ef, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_ef, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_ef, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                n_art = row[materias_cols.get("Educación Artística", 10)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO ARTÍSTICO",
                    "asignaturas": [{
                        "nombre": "EDUCACIÓN ARTÍSTICA",
                        "p1": generar_detalles_etapa(variante_desempeno(n_art, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_art, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_art, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                n_et = row[materias_cols.get("Ética y Valores", 11)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO ÉTICO Y CIUDADANO",
                    "asignaturas": [{
                        "nombre": "EDUCACIÓN ÉTICA Y VALORES",
                        "p1": generar_detalles_etapa(variante_desempeno(n_et, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_et, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_et, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

                n_tec = row[materias_cols.get("Tecnología e Informática", 8)] or "ALTO"
                areas.append({
                    "nombre": "PENSAMIENTO TECNOLÓGICO",
                    "asignaturas": [{
                        "nombre": "TECNOLOGÍA E INFORMÁTICA",
                        "p1": generar_detalles_etapa(variante_desempeno(n_tec, rnd), rnd.randint(1, 10000)),
                        "p2": generar_detalles_etapa(variante_desempeno(n_tec, rnd), rnd.randint(1, 10000)),
                        "p3": generar_detalles_etapa(n_tec, rnd.randint(1, 10000)),
                        "p4": {"p": "-", "rd": "-", "r": "-", "a": "-", "pt": "-", "total": 0, "des": "-"}
                    }]
                })

            convivencia_data = {
                "convivencia": {
                    "p1": {"pa": f"{rnd.randint(78, 95)/10:.1f}".replace(".", ","), "vc": str(rnd.randint(8, 10)), "pc": f"{rnd.randint(78, 95)},{rnd.randint(0, 9)}", "des": rnd.choice(["S", "A", "BS"])},
                    "p2": {"pa": f"{rnd.randint(78, 95)/10:.1f}".replace(".", ","), "vc": str(rnd.randint(8, 10)), "pc": f"{rnd.randint(78, 95)},{rnd.randint(0, 9)}", "des": rnd.choice(["S", "A", "BS"])},
                    "p3": {"pa": f"{rnd.randint(80, 96)/10:.1f}".replace(".", ","), "vc": str(rnd.randint(8, 10)), "pc": f"{rnd.randint(80, 96)},{rnd.randint(0, 9)}", "des": rnd.choice(["S", "A"])},
                    "p4": {"pa": "0,0", "vc": "-", "pc": "0,0", "des": "-"}
                },
                "ambito_familiar": {"p1": "S", "p2": "S", "p3": "S", "p4": "-"},
                "fallas_retardos": {
                    "p1": {"f": rnd.randint(0, 2), "r": rnd.randint(0, 3), "e": 0, "u": 0},
                    "p2": {"f": rnd.randint(0, 3), "r": rnd.randint(0, 3), "e": 0, "u": rnd.randint(0, 1)},
                    "p3": {"f": rnd.randint(0, 1), "r": rnd.randint(0, 2), "e": 0, "u": 0},
                    "p4": {"f": 0, "r": 0, "e": 0, "u": 0}
                }
            }

            datos_academicos = {
                "areas": areas,
                "promedio_actitudinal": {
                    "p1": f"{rnd.randint(78, 95)/10:.2f}".replace(".", ","),
                    "p2": f"{rnd.randint(78, 95)/10:.2f}".replace(".", ","),
                    "p3": f"{rnd.randint(80, 96)/10:.2f}".replace(".", ","),
                    "p4": "0,00"
                },
                "tecnico_sena": {
                    "p1": rnd.choice(["A", "A", "BS"]),
                    "p2": rnd.choice(["A", "BS"]),
                    "p3": rnd.choice(["A", "A", "S"]),
                    "p4": "-"
                }
            }

            datos_por_doc[documento] = (datos_academicos, convivencia_data)

    print(f"Insertando en bloque {len(estudiantes_objs)} estudiantes...")
    Estudiante.objects.bulk_create(
        estudiantes_objs,
        batch_size=50,
        update_conflicts=True,
        unique_fields=["documento"],
        update_fields=["nombres", "apellidos", "grado", "curso", "director_grupo", "acudiente", "activo"],
    )

    mapa_est = {e.documento: e for e in Estudiante.objects.filter(documento__in=list(datos_por_doc.keys()))}
    informes_objs = []
    for doc, (d_acad, d_conv) in datos_por_doc.items():
        est = mapa_est.get(doc)
        if not est:
            continue
        informes_objs.append(
            InformeValorativo(
                estudiante=est,
                anio_lectivo=2026,
                periodo_actual=3,
                datos_academicos=d_acad,
                datos_convivencia=d_conv,
                observaciones_academicas=(
                    "El estudiante demuestra compromiso constante en el desarrollo de las guías SERC.\n"
                    "Se sugiere continuar fortaleciendo la fase de Relación y entrega puntual de evidencias."
                ),
                observaciones_convivenciales=(
                    "La institución agradece su colaboración en las campañas de mejora.\n"
                    "Felicitaciones por su aporte a la convivencia institucional."
                ),
                tecnico_sena=d_acad["tecnico_sena"]["p3"],
            )
        )

    print(f"Insertando en bloque {len(informes_objs)} informes valorativos...")
    InformeValorativo.objects.bulk_create(
        informes_objs,
        batch_size=50,
        update_conflicts=True,
        unique_fields=["estudiante", "anio_lectivo"],
        update_fields=["periodo_actual", "datos_academicos", "datos_convivencia", "observaciones_academicas", "observaciones_convivenciales", "tecnico_sena"],
    )
    print("¡Importación masiva finalizada con éxito!")


if __name__ == "__main__":
    importar()
