import openpyxl

def excel_to_html(wb_path, sheet_name, max_rows=16):
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ws = wb[sheet_name]
    
    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; background: #f1f5f9; padding: 24px; margin: 0; }}
  .excel-window {{ background: #fff; border-radius: 8px; box-shadow: 0 10px 25px rgba(0,0,0,0.1); border: 1px solid #cbd5e1; overflow: hidden; max-width: 1380px; margin: 0 auto; }}
  .excel-titlebar {{ background: #0a1931; color: #fff; padding: 12px 20px; display: flex; align-items: center; justify-content: space-between; font-weight: 700; font-size: 0.95rem; }}
  .tabs {{ display: flex; gap: 4px; background: #e2e8f0; padding: 6px 12px 0; border-bottom: 1px solid #cbd5e1; overflow-x: auto; }}
  .tab {{ padding: 6px 14px; font-size: 0.8rem; font-weight: 600; color: #475569; background: #cbd5e1; border-radius: 4px 4px 0 0; }}
  .tab.active {{ background: #fff; color: #0a1931; border-top: 2px solid #c5a880; }}
  .table-container {{ overflow-x: auto; max-height: 80vh; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 0.78rem; text-align: center; }}
  th {{ background: #0a1931; color: #fff; padding: 10px 8px; font-weight: 700; border: 1px solid #334155; position: sticky; top: 0; white-space: nowrap; }}
  th.subj {{ background: #15305b; }}
  td {{ padding: 8px 6px; border: 1px solid #e2e8f0; white-space: nowrap; }}
  td.num {{ font-weight: 700; color: #64748b; background: #f8fafc; }}
  td.name {{ text-align: left; font-weight: 600; padding-left: 12px; color: #0f172a; }}
  td.doc {{ font-family: monospace; font-size: 0.8rem; color: #334155; }}
  .sup {{ background: #d1e7dd; color: #0f5132; font-weight: 700; border-radius: 3px; padding: 2px 6px; }}
  .alt {{ background: #e8f0fe; color: #155724; font-weight: 700; border-radius: 3px; padding: 2px 6px; }}
  .bas {{ background: #fff3cd; color: #664d03; font-weight: 700; border-radius: 3px; padding: 2px 6px; }}
  .baj {{ background: #f8d7da; color: #842029; font-weight: 700; border-radius: 3px; padding: 2px 6px; }}
</style>
</head>
<body>
  <div class="excel-window">
    <div class="excel-titlebar">
      <div>📊 SÁBANA DE NOTAS CEIS 2026 — I.E.D. Complejo Educativo Integral Sopó</div>
      <div style="font-size: 0.8rem; font-weight: 500; color: #c5a880;">Formato Oficial Decreto 1290 (16 Cursos)</div>
    </div>
    <div class="tabs">
      {''.join(f'<div class="tab {"active" if name == sheet_name else ""}">{name}</div>' for name in wb.sheetnames)}
    </div>
    <div class="table-container">
      <table>
        <thead>
          <tr>
"""
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    for i, h in enumerate(headers, 1):
        cls = "subj" if 4 <= i < len(headers) - 1 else ""
        html += f'            <th class="{cls}">{h}</th>\n'
    html += "          </tr>\n        </thead>\n        <tbody>\n"

    for r in range(2, min(max_rows + 2, ws.max_row + 1)):
        html += "          <tr>\n"
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value or ""
            val_str = str(val)
            if c == 1:
                html += f'            <td class="num">{val_str}</td>\n'
            elif c == 2:
                html += f'            <td class="doc">{val_str}</td>\n'
            elif c == 3:
                html += f'            <td class="name">{val_str}</td>\n'
            elif val_str == "SUPERIOR" or "Excelente" in val_str:
                html += f'            <td><span class="sup">{val_str}</span></td>\n'
            elif val_str == "ALTO" or "Sobresaliente" in val_str:
                html += f'            <td><span class="alt">{val_str}</span></td>\n'
            elif val_str == "BÁSICO" or "Aprobado" in val_str:
                html += f'            <td><span class="bas">{val_str}</span></td>\n'
            elif val_str == "BAJO" or "Requiere" in val_str or "Pendiente" in val_str:
                html += f'            <td><span class="baj">{val_str}</span></td>\n'
            else:
                html += f'            <td>{val_str}</td>\n'
        html += "          </tr>\n"

    html += """        </tbody>
      </table>
    </div>
  </div>
</body>
</html>"""
    return html

if __name__ == "__main__":
    wb_file = "/home/juan/proyectos/athena/sabana_notas_ceis_2026.xlsx"
    with open("/home/juan/proyectos/athena/capturas_exposicion/preview_601.html", "w", encoding="utf-8") as f:
        f.write(excel_to_html(wb_file, "601", max_rows=14))
    with open("/home/juan/proyectos/athena/capturas_exposicion/preview_1001.html", "w", encoding="utf-8") as f:
        f.write(excel_to_html(wb_file, "1001", max_rows=14))
    print("HTML previews generados")
