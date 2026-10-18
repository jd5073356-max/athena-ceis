import os
import random
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Lista de estudiantes reales de 6° si están disponibles
def obtener_estudiantes_sexto():
    estudiantes_por_curso = {"601": [], "602": [], "603": []}
    archivo_original = "/home/juan/Descargas/estudiantes sexto ceis 2026.xlsx"
    if os.path.exists(archivo_original):
        wb = openpyxl.load_workbook(archivo_original, data_only=True)
        for curso in ["601", "602", "603"]:
            if curso in wb.sheetnames:
                ws = wb[curso]
                for r in range(2, ws.max_row + 1):
                    nom = ws.cell(r, 3).value
                    doc = ws.cell(r, 6).value
                    if nom and str(nom).strip():
                        nom_limpio = str(nom).split(" - ")[0].strip().upper()
                        doc_str = ""
                        if doc is not None:
                            try:
                                doc_str = str(int(float(doc)))
                            except Exception:
                                doc_str = str(doc).strip()
                        if not doc_str or len(doc_str) < 6:
                            doc_str = f"1075{random.randint(100000, 999999)}"
                        estudiantes_por_curso[curso].append({
                            "nombre": nom_limpio,
                            "documento": doc_str
                        })
    return estudiantes_por_curso

APELLIDOS = [
    "ACOSTA", "ALARCON", "ALVARADO", "ARANGUREN", "AVILA", "BAQUERO", "BARRETO", "BELTRAN",
    "BERNAL", "BLANCO", "BUITRAGO", "CALDERON", "CAMACHO", "CARDENAS", "CASTIBLANCO", "CASTILLO",
    "CASTRO", "CELIS", "CIFUENTES", "CONTRERAS", "CORTES", "CRUZ", "CUBIDES", "DIAZ",
    "DUARTE", "ESPINOSA", "FORERO", "GAITAN", "GARAVITO", "GARCIA", "GARZON", "GOMEZ",
    "GONZALEZ", "GUACANEME", "GUERRERO", "GUTIERREZ", "HERNANDEZ", "HERRERA", "JIMENEZ", "LARA",
    "LEON", "LOPEZ", "LOZANO", "MARIN", "MARTINEZ", "MEDINA", "MEJIA", "MOLINA", "MONCADA",
    "MORALES", "MORENO", "MUNAR", "MUÑOZ", "NAVARRO", "NARIÑO", "NOVOA", "ORTEGA", "ORTIZ",
    "OSPINA", "PARRA", "PENA", "PEREZ", "PIRAQUIVE", "PULIDO", "QUINTERO", "RAMIREZ", "RAMOS",
    "RINCON", "RIVERA", "RODRIGUEZ", "ROJAS", "ROMERO", "RUIZ", "SAENZ", "SALAZAR", "SANCHEZ",
    "SANDOVAL", "SIERRA", "SILVA", "SUAREZ", "TORRES", "TRUJILLO", "URIBE", "VALENCIA",
    "VARGAS", "VEGA", "VELASQUEZ", "VERGARA", "VILLAMIL", "VILLAMIZAR", "ZAMORA", "ZAPATA"
]

NOMBRES_MASC = [
    "ALEJANDRO", "ANDRES", "CAMILO", "CRISTIAN", "DANIEL", "DAVID", "DIEGO", "DYLAN",
    "EDER", "ESTEBAN", "FELIPE", "GABRIEL", "IAN", "ITAN", "JACOBO", "JAVIER", "JERONIMO",
    "JHOAN", "JOAQUIN", "JORGE", "JOSE", "JUAN", "JULIAN", "MANUEL", "MARTIN", "MATEO",
    "MATHIAS", "MIGUEL", "NICOLAS", "SAMUEL", "SANTIAGO", "SEBASTIAN", "TOMAS", "VALENTIN"
]

NOMBRES_FEM = [
    "ANA", "ANDREA", "ANTONIA", "CAMILA", "DANIELA", "DULCE", "EMMA", "GABRIELA",
    "ISABELLA", "JULIANA", "LAURA", "LUCIANA", "LUISA", "MANUELA", "MARIA", "MARIANA",
    "PAULA", "SALOME", "SARA", "SARAY", "SOFIA", "VALENTINA", "VALERIA", "VICTORIA"
]

def generar_estudiantes_aleatorios(cantidad, base_doc=1075000000):
    lista = []
    usados = set()
    while len(lista) < cantidad:
        ap1 = random.choice(APELLIDOS)
        ap2 = random.choice(APELLIDOS)
        if random.random() < 0.5:
            n1 = random.choice(NOMBRES_MASC)
            n2 = random.choice(NOMBRES_MASC)
        else:
            n1 = random.choice(NOMBRES_FEM)
            n2 = random.choice(NOMBRES_FEM)
        nombre_completo = f"{ap1} {ap2} {n1} {n2}"
        if nombre_completo not in usados:
            usados.add(nombre_completo)
            doc_num = base_doc + random.randint(10000, 999999)
            lista.append({
                "nombre": nombre_completo,
                "documento": str(doc_num)
            })
    lista.sort(key=lambda x: x["nombre"])
    return lista

# Materias oficiales según requerimiento
MATERIAS_SECUNDARIA = [
    "Matemáticas",
    "Lengua Castellana",
    "Ciencias Naturales",
    "Ciencias Sociales",
    "Inglés",
    "Tecnología e Informática",
    "Educación Física",
    "Educación Artística",
    "Ética y Valores"
]

MATERIAS_MEDIA = [
    "Matemáticas",
    "Lengua Castellana",
    "Física",
    "Química",
    "Política y Economía",
    "Filosofía",
    "Inglés",
    "Tecnología e Informática",
    "Educación Física",
    "Educación Artística",
    "Ética y Valores"
]

CURSOS = [
    "601", "602", "603",
    "701", "702", "703",
    "801", "802", "803",
    "901", "902", "903",
    "1001", "1002",
    "1101", "1102"
]

# Estilos visuales
NAVY_HEADER = PatternFill(start_color="0A1931", end_color="0A1931", fill_type="solid")
GOLD_HEADER = PatternFill(start_color="15305B", end_color="15305B", fill_type="solid")
WHITE_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="0A1931")
REGULAR_FONT = Font(name="Calibri", size=10, color="1E293B")

FILL_SUPERIOR = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")  # Verde menta suave
FONT_SUPERIOR = Font(name="Calibri", size=10, bold=True, color="0F5132")

FILL_ALTO = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")      # Azul suave
FONT_ALTO = Font(name="Calibri", size=10, bold=True, color="155724")

FILL_BASICO = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")    # Amarillo/ámbar suave
FONT_BASICO = Font(name="Calibri", size=10, bold=True, color="664D03")

FILL_BAJO = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")      # Rojo suave
FONT_BAJO = Font(name="Calibri", size=10, bold=True, color="842029")

THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")

def generar_nota_cualitativa():
    # Distribución realista de calificaciones escolares
    r = random.random()
    if r < 0.28:
        return "SUPERIOR"
    elif r < 0.72:
        return "ALTO"
    elif r < 0.94:
        return "BÁSICO"
    else:
        return "BAJO"

def calcular_desempeno_general(notas):
    bajos = notas.count("BAJO")
    if bajos >= 2:
        return "BAJO", "Requiere Plan de Apoyo"
    elif bajos == 1:
        return "BÁSICO", "1 Asignatura Pendiente"
    
    superiores = notas.count("SUPERIOR")
    altos = notas.count("ALTO")
    if superiores >= len(notas) * 0.6:
        return "SUPERIOR", "Al Día (Excelente)"
    elif (superiores + altos) >= len(notas) * 0.7:
        return "ALTO", "Al Día (Sobresaliente)"
    else:
        return "BÁSICO", "Al Día (Aprobado)"

def crear_libro_excel(ruta_salida):
    wb = openpyxl.Workbook()
    # Eliminar hoja por defecto
    default_sheet = wb.active

    estudiantes_sexto = obtener_estudiantes_sexto()

    for idx, curso in enumerate(CURSOS):
        ws = wb.create_sheet(title=curso)
        es_media = curso in ["1001", "1002", "1101", "1102"]
        materias = MATERIAS_MEDIA if es_media else MATERIAS_SECUNDARIA

        # Encabezados
        headers = ["No.", "Documento", "Estudiante"] + materias + ["Desempeño General", "Estado Académico"]
        ws.append(headers)

        # Aplicar estilo al encabezado
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = NAVY_HEADER if col_idx <= 3 else GOLD_HEADER
            cell.font = WHITE_FONT
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        
        ws.row_dimensions[1].height = 28

        # Obtener lista de estudiantes para este curso
        if curso in estudiantes_sexto and len(estudiantes_sexto[curso]) > 0:
            estudiantes = estudiantes_sexto[curso]
            estudiantes.sort(key=lambda x: x["nombre"])
        else:
            cant = 32 if es_media else random.randint(32, 36)
            base_doc = 1075000000 + (int(curso) * 10000)
            estudiantes = generar_estudiantes_aleatorios(cant, base_doc)

        # Rellenar filas de estudiantes
        for row_num, est in enumerate(estudiantes, start=2):
            notas_fila = [generar_nota_cualitativa() for _ in materias]
            desempeno_gen, estado = calcular_desempeno_general(notas_fila)

            row_data = [row_num - 1, est["documento"], est["nombre"]] + notas_fila + [desempeno_gen, estado]
            ws.append(row_data)
            ws.row_dimensions[row_num].height = 20

            # Estilos de celdas
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.border = THIN_BORDER
                cell.font = REGULAR_FONT
                
                if col_idx == 1:
                    cell.alignment = ALIGN_CENTER
                    cell.font = BOLD_FONT
                elif col_idx == 2:
                    cell.alignment = ALIGN_CENTER
                elif col_idx == 3:
                    cell.alignment = ALIGN_LEFT
                    cell.font = BOLD_FONT
                elif 4 <= col_idx < 4 + len(materias):
                    cell.alignment = ALIGN_CENTER
                    if val == "SUPERIOR":
                        cell.fill = FILL_SUPERIOR
                        cell.font = FONT_SUPERIOR
                    elif val == "ALTO":
                        cell.fill = FILL_ALTO
                        cell.font = FONT_ALTO
                    elif val == "BÁSICO":
                        cell.fill = FILL_BASICO
                        cell.font = FONT_BASICO
                    elif val == "BAJO":
                        cell.fill = FILL_BAJO
                        cell.font = FONT_BAJO
                else:
                    # Desempeño y Estado
                    cell.alignment = ALIGN_CENTER
                    cell.font = BOLD_FONT
                    if val in ["SUPERIOR", "Al Día (Excelente)"]:
                        cell.fill = FILL_SUPERIOR
                        cell.font = FONT_SUPERIOR
                    elif val in ["ALTO", "Al Día (Sobresaliente)"]:
                        cell.fill = FILL_ALTO
                        cell.font = FONT_ALTO
                    elif val in ["BÁSICO", "Al Día (Aprobado)"]:
                        cell.fill = FILL_BASICO
                        cell.font = FONT_BASICO
                    elif "BAJO" in val or "Requiere" in val or "Pendiente" in val:
                        cell.fill = FILL_BAJO
                        cell.font = FONT_BAJO

        # Fijar panel superior para scroll cómodo
        ws.freeze_panes = "D2"

        # Autoajustar ancho de columnas
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)
        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 38

    # Remover default sheet
    wb.remove(default_sheet)

    wb.save(ruta_salida)
    print(f"Libro guardado exitosamente en: {ruta_salida}")

if __name__ == "__main__":
    salida = "/home/juan/proyectos/athena/sabana_notas_ceis_2026.xlsx"
    crear_libro_excel(salida)
    # También copiar a Descargas para fácil acceso
    import shutil
    shutil.copyfile(salida, "/home/juan/Descargas/sabana_notas_ceis_2026.xlsx")
    print("Copia creada en /home/juan/Descargas/sabana_notas_ceis_2026.xlsx")
